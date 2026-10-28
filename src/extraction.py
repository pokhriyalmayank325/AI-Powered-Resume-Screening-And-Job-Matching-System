"""Step 1: PDF / DOCX / TXT  ->  plain text."""
from __future__ import annotations

import io
from pathlib import Path


def _read_bytes(source, filename: str | None):
    """Accept a path, raw bytes, or a Streamlit UploadedFile."""
    if isinstance(source, (str, Path)):
        p = Path(source)
        return p.read_bytes(), p.name
    if isinstance(source, (bytes, bytearray)):
        return bytes(source), filename or ""
    data = source.read()  # file-like (Streamlit UploadedFile)
    if hasattr(source, "seek"):
        source.seek(0)
    return data, filename or getattr(source, "name", "")


def _pdf_to_text(data: bytes) -> str:
    import fitz  # PyMuPDF

    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text("text") for page in doc)


def _docx_to_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:  # many resumes put skills in tables
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(source, filename: str | None = None) -> str:
    data, name = _read_bytes(source, filename)
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return _pdf_to_text(data)
    if suffix == ".docx":
        return _docx_to_text(data)
    if suffix in {".txt", ".md", ""}:
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {suffix or 'unknown'} (use PDF, DOCX or TXT)")
